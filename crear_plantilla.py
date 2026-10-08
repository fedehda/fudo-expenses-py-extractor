#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Genera la Plantilla-Gastos.xlsx oficial de Fudo con sus 2 hojas:
1. Gastos (11 columnas)
2. Listas (Categorías, Proveedores, Tipo de comprobante, Medio de pago, Caja)
"""
import openpyxl
from pathlib import Path

def crear_plantilla_fudo(ruta_destino: Path):
    wb = openpyxl.Workbook()

    # Hoja 1: Gastos
    ws_gastos = wb.active
    ws_gastos.title = "Gastos"
    columnas_gastos = [
        "Fecha",
        "Monto",
        "Proveedor",
        "Categoría",
        "Comentario",
        "Tipo de comprobante",
        "Número de comprobante",
        "Medio de pago",
        "Caja",
        "Fecha de pago",
        "Fecha de vencimiento",
    ]
    for col_idx, col_name in enumerate(columnas_gastos, start=1):
        ws_gastos.cell(row=1, column=col_idx, value=col_name)

    # Hoja 2: Listas
    ws_listas = wb.create_sheet(title="Listas")

    categorias = [
        "Compras y proveedores", "Carnicería", "Verdulería", "Almacén", "Bebidas",
        "Panadería", "Limpieza y mantenimiento", "Servicios", "Impuestos",
        "Alquileres", "Sueldos y jornales", "Gastos de librería y oficina",
        "Publicidad y marketing", "Mantenimiento y reparaciones", "Comisiones",
        "Seguros", "Honorarios profesionales", "Combustibles y viáticos",
        "Uniformes y ropa de trabajo", "Equipamiento y vajilla", "Flete y traslados",
        "Seguridad y vigilancia", "Capacitación y cursos", "Telefonía e internet",
        "Energía eléctrica", "Gas natural / envasado", "Agua y saneamiento",
        "Retenciones bancarias", "Gastos financieros", "Otros gastos"
    ]

    proveedores = [
        "Proveedor Ejemplo A",
        "Proveedor Ejemplo B",
        "Distribuidora Ejemplo SRL",
    ]

    tipos_comprobante = ["Factura A", "Factura B", "Factura C", "Recibo", "Remito"]
    medios_pago = ["Efectivo", "Cta. Cte.", "Payway", "Transferencia", "MercadoPago"]
    cajas = ["Principal"]

    # Escribir encabezados
    ws_listas.cell(row=1, column=1, value="Categoría")
    ws_listas.cell(row=1, column=2, value="Proveedor")
    ws_listas.cell(row=1, column=3, value="Tipo de comprobante")
    ws_listas.cell(row=1, column=4, value="Medio de pago")
    ws_listas.cell(row=1, column=5, value="Caja")

    for r, v in enumerate(categorias, start=2):
        ws_listas.cell(row=r, column=1, value=v)

    for r, v in enumerate(proveedores, start=2):
        ws_listas.cell(row=r, column=2, value=v)

    for r, v in enumerate(tipos_comprobante, start=2):
        ws_listas.cell(row=r, column=3, value=v)

    for r, v in enumerate(medios_pago, start=2):
        ws_listas.cell(row=r, column=4, value=v)

    for r, v in enumerate(cajas, start=2):
        ws_listas.cell(row=r, column=5, value=v)

    wb.save(ruta_destino)
    print(f"[OK] Plantilla creada exitosamente en {ruta_destino} con {len(categorias)} categorias y {len(proveedores)} proveedores.")


if __name__ == "__main__":
    destino = Path(__file__).resolve().parent / "Plantilla-Gastos.xlsx"
    crear_plantilla_fudo(destino)
