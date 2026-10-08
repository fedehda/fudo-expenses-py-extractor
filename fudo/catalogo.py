"""Catálogos de la plantilla de Fudo (hoja Listas) y coincidencia aproximada."""
import difflib
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import openpyxl

from .config import BASE_DIR
from .validaciones import normalizar_texto

CATEGORIAS_BASE = [
    "Compras y proveedores", "Carnicería", "Verdulería", "Almacén", "Bebidas", "Panadería",
    "Limpieza y mantenimiento", "Servicios", "Impuestos", "Alquileres", "Sueldos y jornales",
    "Gastos de librería y oficina", "Publicidad y marketing", "Mantenimiento y reparaciones",
    "Comisiones", "Otros gastos",
]


def mejor_coincidencia(nombre: Optional[str], opciones: List[str]) -> Tuple[Optional[str], float]:
    """Devuelve (opción, similitud 0..1) comparando textos normalizados."""
    if not nombre or not opciones:
        return None, 0.0
    objetivo = normalizar_texto(nombre)
    if not objetivo:
        return None, 0.0
    mejor, puntaje = None, 0.0
    for op in opciones:
        n = normalizar_texto(op)
        if not n:
            continue
        if n == objetivo:
            return op, 1.0
        r = difflib.SequenceMatcher(None, objetivo, n).ratio()
        # Un nombre contenido en el otro ("la serenisima" vs "mastellone la serenisima")
        if len(n) >= 5 and len(objetivo) >= 5 and (n in objetivo or objetivo in n):
            r = max(r, 0.9)
        if r > puntaje:
            mejor, puntaje = op, r
    return mejor, puntaje


class CatalogoFudo:
    def __init__(self, template_path: Optional[Path] = None):
        self.categorias: List[str] = []
        self.proveedores: List[str] = []
        self.tipos_comprobante = ["Factura A", "Factura B", "Factura C", "Recibo", "Remito"]
        self.medios_pago = ["Efectivo", "Cta. Cte.", "Payway", "Transferencia", "MercadoPago"]
        self.cajas = ["Principal"]
        self.template_file: Optional[Path] = None
        self._cargar(template_path)
        if not self.categorias:
            self.categorias = list(CATEGORIAS_BASE)

    def _cargar(self, template_path: Optional[Path]):
        for c in (template_path, BASE_DIR / "Plantilla-Gastos.xlsx", BASE_DIR / "Plantilla-Gastos (1).xlsx"):
            if c and Path(c).exists():
                self.template_file = Path(c)
                break
        if not self.template_file:
            return
        try:
            wb = openpyxl.load_workbook(self.template_file, data_only=True)
        except Exception as e:
            print(f"[!] No se pudo leer {self.template_file.name}: {e}. Se usan catálogos base.")
            return
        try:
            if "Listas" not in wb.sheetnames:
                return
            ws = wb["Listas"]
            headers: Dict[str, int] = {}
            for col in range(1, ws.max_column + 1):
                v = ws.cell(row=1, column=col).value
                if v:
                    headers[normalizar_texto(v)] = col

            def columna(nombres: List[str]) -> List[str]:
                col = next((headers[normalizar_texto(n)] for n in nombres if normalizar_texto(n) in headers), None)
                if not col:
                    return []
                return [str(ws.cell(row=r, column=col).value).strip() for r in range(2, ws.max_row + 1)
                        if ws.cell(row=r, column=col).value is not None and str(ws.cell(row=r, column=col).value).strip()]

            self.categorias = columna(["Categoría", "Categorias"]) or self.categorias
            self.proveedores = columna(["Proveedor", "Proveedores"]) or self.proveedores
            self.tipos_comprobante = columna(["Tipo de comprobante", "Comprobante"]) or self.tipos_comprobante
            self.medios_pago = columna(["Medio de pago", "Forma de pago"]) or self.medios_pago
            self.cajas = columna(["Caja", "Cajas"]) or self.cajas
        finally:
            wb.close()

    def categoria_valida(self, nombre: Optional[str]) -> Tuple[Optional[str], float]:
        return mejor_coincidencia(nombre, self.categorias)

    def proveedor_catalogo(self, nombre: Optional[str]) -> Tuple[Optional[str], float]:
        return mejor_coincidencia(nombre, self.proveedores)
