"""Módulo analítico de Histórico de Precios, Comparador de Proveedores y Alertas."""
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .validaciones import parsear_fecha


def normalizar_producto(texto: Optional[str]) -> str:
    """Normaliza la descripción de un artículo para agruparlo y compararlo limpiamente."""
    if not texto:
        return ""
    # Quitar tildes y caracteres especiales
    t = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode("ascii").lower()
    # Reemplazar símbolos y puntuación por espacios
    t = re.sub(r"[^a-z0-9 ]+", " ", t)
    # Quitar palabras de embalaje o unidades típicas sueltas
    t = re.sub(r"\b(cajon|cajones|bolsa|bolsas|pack|packs|bulto|bultos|unidades?|unid?|kgr?|lts?|gr?)\b", " ", t)
    # Quitar 'x' aislada o multiplicadores como 'x 10'
    t = re.sub(r"\b(x\s*\d+|x)\b", " ", t)
    t = re.sub(r"\s+", " ", t).strip()
    return t or str(texto).strip().lower()


def inferir_precio_unitario(importe: Optional[float], cantidad_num: Optional[float],
                            precio_unit_existente: Optional[float]) -> Optional[float]:
    """Calcula o valida el precio unitario de una línea."""
    if precio_unit_existente is not None and precio_unit_existente > 0:
        return round(float(precio_unit_existente), 2)
    if importe is not None and importe > 0 and cantidad_num is not None and cantidad_num > 0:
        return round(float(importe) / float(cantidad_num), 2)
    if importe is not None and importe > 0:
        return round(float(importe), 2)
    return None


@dataclass
class ItemHistorico:
    id: int
    comprobante_hash: str
    fecha: str
    proveedor_id: Optional[int]
    proveedor_nombre: str
    producto_original: str
    producto_normalizado: str
    cantidad_texto: Optional[str]
    cantidad_numerica: Optional[float]
    unidad_medida: Optional[str]
    precio_unitario: Optional[float]
    importe_total: Optional[float]
    tipo_comprobante: Optional[str]
    numero_comprobante: Optional[str]
    archivo: Optional[str]


@dataclass
class ResumenProducto:
    producto_normalizado: str
    nombre_mostrar: str
    ultimo_precio: float
    ultima_fecha: str
    ultimo_proveedor: str
    precio_anterior: Optional[float] = None
    fecha_anterior: Optional[str] = None
    variacion_pct: float = 0.0
    mejor_precio: float = 0.0
    mejor_proveedor: str = ""
    peor_precio: float = 0.0
    peor_proveedor: str = ""
    total_compras: int = 0
    proveedores: List[str] = field(default_factory=list)


def calcular_resumen_producto(items: List[ItemHistorico]) -> Optional[ResumenProducto]:
    """A partir de la lista cronológica de compras de un producto, genera el análisis de precios."""
    if not items:
        return None

    # Ordenar por fecha cronológica (más antiguo a más reciente)
    def clave_fecha(it: ItemHistorico):
        d = parsear_fecha(it.fecha)
        return (d if d else datetime.min.date(), it.id)

    items_ordenados = sorted(items, key=clave_fecha)
    # Filtrar aquellos con precio unitario válido
    con_precio = [it for it in items_ordenados if it.precio_unitario is not None and it.precio_unitario > 0]
    if not con_precio:
        con_precio = items_ordenados

    ultimo = con_precio[-1]
    anterior = con_precio[-2] if len(con_precio) >= 2 else None

    p_ultimo = ultimo.precio_unitario or 0.0
    p_anterior = anterior.precio_unitario if anterior else None

    var_pct = 0.0
    if p_anterior and p_anterior > 0:
        var_pct = round(((p_ultimo - p_anterior) / p_anterior) * 100.0, 1)

    # Determinar mejor y peor proveedor por precio unitario
    mejor_item = min(con_precio, key=lambda x: x.precio_unitario or float("inf"))
    peor_item = max(con_precio, key=lambda x: x.precio_unitario or 0.0)

    provs = sorted(list(dict.fromkeys(it.proveedor_nombre for it in items if it.proveedor_nombre)))

    # Tomar el nombre original más descriptivo o más reciente
    nombre_mostrar = ultimo.producto_original or items[0].producto_original

    return ResumenProducto(
        producto_normalizado=items[0].producto_normalizado,
        nombre_mostrar=nombre_mostrar,
        ultimo_precio=p_ultimo,
        ultima_fecha=ultimo.fecha,
        ultimo_proveedor=ultimo.proveedor_nombre,
        precio_anterior=p_anterior,
        fecha_anterior=anterior.fecha if anterior else None,
        variacion_pct=var_pct,
        mejor_precio=mejor_item.precio_unitario or p_ultimo,
        mejor_proveedor=mejor_item.proveedor_nombre,
        peor_precio=peor_item.precio_unitario or p_ultimo,
        peor_proveedor=peor_item.proveedor_nombre,
        total_compras=len(items),
        proveedores=provs,
    )


def exportar_historico_precios_excel(items_todos: List[ItemHistorico],
                                     resumenes: List[ResumenProducto],
                                     ruta_destino: Path):
    """Genera un libro de Excel analítico con resumen de precios, alertas e historial."""
    wb = openpyxl.Workbook()
    # -------------------------------------------------------------
    # Hoja 1: Resumen y Comparador
    # -------------------------------------------------------------
    ws_resumen = wb.active
    ws_resumen.title = "Resumen y Comparador"

    headers_resumen = [
        "Producto", "Mejor Proveedor", "Mejor Precio Unit.",
        "Último Proveedor", "Último Precio Unit.", "Precio Anterior",
        "Variación %", "Última Compra", "Total Compras", "Proveedores Distintos"
    ]

    header_fill = PatternFill(start_color="1B4F72", end_color="1B4F72", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    border_fino = Border(
        left=Side(style="thin", color="D0D3D4"),
        right=Side(style="thin", color="D0D3D4"),
        top=Side(style="thin", color="D0D3D4"),
        bottom=Side(style="thin", color="D0D3D4"),
    )

    for col_idx, h in enumerate(headers_resumen, start=1):
        c = ws_resumen.cell(row=1, column=col_idx, value=h)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center" if col_idx > 1 else "left", vertical="center")

    fill_rojo = PatternFill(start_color="FADBD8", end_color="FADBD8", fill_type="solid")
    fill_verde = PatternFill(start_color="D4EFDF", end_color="D4EFDF", fill_type="solid")

    for r_idx, r in enumerate(resumenes, start=2):
        ws_resumen.cell(row=r_idx, column=1, value=r.nombre_mostrar)
        ws_resumen.cell(row=r_idx, column=2, value=r.mejor_proveedor)
        c_mp = ws_resumen.cell(row=r_idx, column=3, value=r.mejor_precio)
        c_mp.number_format = "$#,##0.00"
        ws_resumen.cell(row=r_idx, column=4, value=r.ultimo_proveedor)
        c_up = ws_resumen.cell(row=r_idx, column=5, value=r.ultimo_precio)
        c_up.number_format = "$#,##0.00"

        c_ant = ws_resumen.cell(row=r_idx, column=6, value=r.precio_anterior if r.precio_anterior else "-")
        if r.precio_anterior:
            c_ant.number_format = "$#,##0.00"

        c_var = ws_resumen.cell(row=r_idx, column=7, value=f"{r.variacion_pct:+.1f}%")
        c_var.alignment = Alignment(horizontal="center")
        if r.variacion_pct >= 10.0:
            c_var.fill = fill_rojo
            c_var.font = Font(bold=True, color="922B21")
        elif r.variacion_pct <= -5.0:
            c_var.fill = fill_verde
            c_var.font = Font(bold=True, color="196F3D")

        ws_resumen.cell(row=r_idx, column=8, value=r.ultima_fecha).alignment = Alignment(horizontal="center")
        ws_resumen.cell(row=r_idx, column=9, value=r.total_compras).alignment = Alignment(horizontal="center")
        ws_resumen.cell(row=r_idx, column=10, value=", ".join(r.proveedores))

        for c_i in range(1, len(headers_resumen) + 1):
            ws_resumen.cell(row=r_idx, column=c_i).border = border_fino

    # -------------------------------------------------------------
    # Hoja 2: Alertas de Aumentos
    # -------------------------------------------------------------
    ws_alertas = wb.create_sheet(title="Alertas de Aumento")
    headers_alertas = [
        "Producto", "Proveedor Actual", "Precio Anterior", "Precio Nuevo",
        "Aumento ($)", "Aumento (%)", "Fecha Compra"
    ]
    header_fill_alerta = PatternFill(start_color="922B21", end_color="922B21", fill_type="solid")
    for col_idx, h in enumerate(headers_alertas, start=1):
        c = ws_alertas.cell(row=1, column=col_idx, value=h)
        c.fill = header_fill_alerta
        c.font = header_font
        c.alignment = Alignment(horizontal="center" if col_idx > 1 else "left", vertical="center")

    fila_alerta = 2
    for r in resumenes:
        if r.variacion_pct >= 10.0 and r.precio_anterior:
            dif_pesos = r.ultimo_precio - r.precio_anterior
            ws_alertas.cell(row=fila_alerta, column=1, value=r.nombre_mostrar)
            ws_alertas.cell(row=fila_alerta, column=2, value=r.ultimo_proveedor)
            c1 = ws_alertas.cell(row=fila_alerta, column=3, value=r.precio_anterior)
            c1.number_format = "$#,##0.00"
            c2 = ws_alertas.cell(row=fila_alerta, column=4, value=r.ultimo_precio)
            c2.number_format = "$#,##0.00"
            c3 = ws_alertas.cell(row=fila_alerta, column=5, value=dif_pesos)
            c3.number_format = "$#,##0.00"
            c4 = ws_alertas.cell(row=fila_alerta, column=6, value=f"+{r.variacion_pct:.1f}%")
            c4.fill = fill_rojo
            c4.font = Font(bold=True, color="922B21")
            c4.alignment = Alignment(horizontal="center")
            ws_alertas.cell(row=fila_alerta, column=7, value=r.ultima_fecha).alignment = Alignment(horizontal="center")
            fila_alerta += 1

    # -------------------------------------------------------------
    # Hoja 3: Detalle Histórico
    # -------------------------------------------------------------
    ws_detalle = wb.create_sheet(title="Detalle Histórico")
    headers_detalle = [
        "Fecha", "Proveedor", "Producto Original", "Cantidad",
        "Unidad", "Precio Unitario", "Importe Total", "Tipo Comprobante",
        "Número", "Archivo"
    ]
    header_fill_det = PatternFill(start_color="2E4053", end_color="2E4053", fill_type="solid")
    for col_idx, h in enumerate(headers_detalle, start=1):
        c = ws_detalle.cell(row=1, column=col_idx, value=h)
        c.fill = header_fill_det
        c.font = header_font
        c.alignment = Alignment(horizontal="center" if col_idx in (1, 4, 5, 8, 9) else "left", vertical="center")

    for r_idx, it in enumerate(items_todos, start=2):
        ws_detalle.cell(row=r_idx, column=1, value=it.fecha).alignment = Alignment(horizontal="center")
        ws_detalle.cell(row=r_idx, column=2, value=it.proveedor_nombre)
        ws_detalle.cell(row=r_idx, column=3, value=it.producto_original)
        ws_detalle.cell(row=r_idx, column=4, value=it.cantidad_numerica or it.cantidad_texto or "-")
        ws_detalle.cell(row=r_idx, column=5, value=it.unidad_medida or "-").alignment = Alignment(horizontal="center")
        c_pu = ws_detalle.cell(row=r_idx, column=6, value=it.precio_unitario or 0.0)
        c_pu.number_format = "$#,##0.00"
        c_tot = ws_detalle.cell(row=r_idx, column=7, value=it.importe_total or 0.0)
        c_tot.number_format = "$#,##0.00"
        ws_detalle.cell(row=r_idx, column=8, value=it.tipo_comprobante or "-").alignment = Alignment(horizontal="center")
        ws_detalle.cell(row=r_idx, column=9, value=it.numero_comprobante or "-").alignment = Alignment(horizontal="center")
        ws_detalle.cell(row=r_idx, column=10, value=it.archivo or "-")

    # Ajuste de ancho de columnas en todas las hojas
    for ws in (ws_resumen, ws_alertas, ws_detalle):
        for col in ws.columns:
            longitud = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(longitud + 3, 12)

    ruta_destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta_destino)
    wb.close()
