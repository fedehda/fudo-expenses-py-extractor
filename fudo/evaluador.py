"""Módulo de evaluación y benchmarking con tickets reales.
Ejecuta pruebas sin mover archivos ni alterar la memoria SQLite real.
Compara contra 'pruebas/gabarito.xlsx' o genera un borrador para revisión.
"""
import shutil
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.styles import Font, PatternFill

from .catalogo import CatalogoFudo
from .config import BASE_DIR, Config, EXTENSIONES_VALIDAS
from .extractores import MotorIA
from .memoria import Memoria
from .router import RouterTickets
from .validaciones import formatear_fecha_fudo


import time


def ejecutar_evaluacion(carpeta_pruebas: Optional[Path] = None, modelo: Optional[str] = None,
                        limite: Optional[int] = None, timeout_max_segundos: Optional[int] = None):
    dir_pruebas = Path(carpeta_pruebas or (BASE_DIR / "pruebas" / "tickets_reales")).resolve()
    dir_pruebas.mkdir(parents=True, exist_ok=True)

    # Búsqueda recursiva: encuentra comprobantes tanto en la carpeta raíz como en subcarpetas
    archivos = sorted([p for p in dir_pruebas.rglob("*") if p.is_file() and p.suffix.lower() in EXTENSIONES_VALIDAS])

    print("\n" + "=" * 65)
    print("       EVALUADOR DE PRECISIÓN - TICKETS REALES")
    print("=" * 65)
    print(f"[*] Carpeta de pruebas: {dir_pruebas}")
    print(f"[*] Comprobantes encontrados: {len(archivos)}")

    if not archivos:
        print("\n[i] No hay comprobantes cargados en la carpeta de pruebas.")
        print(f"    Coloca tus tickets reales (JPG, PNG o PDF) en:")
        print(f"    {dir_pruebas}")
        print("\n    Sugerencia de lote:")
        print("    - 8 a 10 Facturas electrónicas con QR")
        print("    - 3 a 5 Facturas sin QR o con QR borroso")
        print("    - 5 a 8 Tiques fiscales de controlador")
        print("    - 3 a 5 Remitos o presupuestos sin nombre del mismo proveedor (para probar memoria)")
        print("=" * 65)
        return

    if limite and limite > 0:
        print(f"[*] Límite de comprobantes a evaluar: {limite} de {len(archivos)}")
        archivos = archivos[:limite]

    if timeout_max_segundos:
        print(f"[*] Límite de tiempo global: {timeout_max_segundos}s")

    # Usar base de datos temporal para no alterar la memoria real durante la evaluación
    with tempfile.TemporaryDirectory() as tmpdir:
        db_temp = Path(tmpdir) / "eval_memoria.db"
        db_real = BASE_DIR / "memoria.db"
        if db_real.exists():
            shutil.copy2(db_real, db_temp)

        memoria_eval = Memoria(db_temp)
        with memoria_eval._conn() as conn_eval:
            conn_eval.execute("DELETE FROM comprobantes")
        catalogo = CatalogoFudo()
        motor_ia = MotorIA(modelo=modelo)
        router = RouterTickets(catalogo=catalogo, memoria=memoria_eval, motor_ia=motor_ia)

        print(f"[*] Motor configurado: {motor_ia.descripcion()}")
        if not motor_ia.disponible:
            print("    [!] AVISO: Ninguna clave de IA activa en .env. Se usará solo lectura QR y OCR local.")
        print("-" * 65)

        resultados = []
        lote_id = f"eval_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
        inicio_global = time.time()

        for idx, arch in enumerate(archivos, start=1):
            if timeout_max_segundos and (time.time() - inicio_global) >= timeout_max_segundos:
                print(f"\n[!] Límite de tiempo global de {timeout_max_segundos}s alcanzado. Deteniendo lote.")
                break

            t0 = time.time()
            print(f"[{idx}/{len(archivos)}] Evaluando: {arch.name}...", end="", flush=True)
            res = router.procesar(arch, lote_id=lote_id)
            resultados.append(res)
            dt = time.time() - t0
            estado_tag = f"[{res.estado}]"
            prov_tag = res.fila_gastos.get("proveedor", "") if res.fila_gastos else "N/A"
            monto_tag = f"${res.fila_gastos.get('monto', 0.0):,.2f}" if res.fila_gastos else "$0.00"
            print(f" ({dt:.1f}s)\n    {estado_tag} {res.estrategia} | Prov: {prov_tag} | {monto_tag}")
            if res.motivos:
                print(f"      Alertas: {'; '.join(res.motivos[:2])}")
            if idx < len(archivos) and Config.pausa_entre_tickets > 0:
                time.sleep(Config.pausa_entre_tickets)

        # Comprobar si existe gabarito de referencia
        ruta_gabarito = BASE_DIR / "pruebas" / "gabarito.xlsx"
        if ruta_gabarito.exists():
            print("\n[*] Comparando contra gabarito oficial (pruebas/gabarito.xlsx)...")
            _comparar_con_gabarito(resultados, ruta_gabarito)
        else:
            print("\n[*] No se encontró 'pruebas/gabarito.xlsx'.")
            ruta_borrador = BASE_DIR / "pruebas" / "gabarito_borrador.xlsx"
            _generar_borrador_gabarito(resultados, ruta_borrador)
            print(f"[OK] Se generó un borrador de referencia en:")
            print(f"     {ruta_borrador}")
            print("     Ábrelo, revisa/corrige los valores y guárdalo como 'gabarito.xlsx'")
            print("     para medir la precisión exacta de cada campo en las siguientes corridas.")

    print("\n[OK] Evaluación finalizada.")
    print("=" * 65)


def _generar_borrador_gabarito(resultados, ruta_destino: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Gabarito"

    headers = [
        "Archivo", "Fecha", "Monto", "Proveedor", "Categoría",
        "Tipo de comprobante", "Número de comprobante", "Medio de pago",
        "CUIT", "Tipo Documento", "Estrategia", "Estado"
    ]
    header_fill = PatternFill(start_color="D6EAF8", end_color="D6EAF8", fill_type="solid")
    header_font = Font(bold=True)

    for col_idx, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col_idx, value=h)
        c.fill = header_fill
        c.font = header_font

    for row_idx, r in enumerate(resultados, start=2):
        fg = r.fila_gastos or {}
        ext = r.extraccion
        cuit_val = ext.cuit_emisor if ext else ""
        tipo_doc_val = ext.tipo_documento.value if ext and hasattr(ext.tipo_documento, "value") else ""

        ws.cell(row=row_idx, column=1, value=r.archivo.name)
        ws.cell(row=row_idx, column=2, value=formatear_fecha_fudo(fg.get("fecha", "")))
        ws.cell(row=row_idx, column=3, value=fg.get("monto", 0.0))
        ws.cell(row=row_idx, column=4, value=fg.get("proveedor", ""))
        ws.cell(row=row_idx, column=5, value=fg.get("categoria", ""))
        ws.cell(row=row_idx, column=6, value=fg.get("tipo_comprobante", ""))
        ws.cell(row=row_idx, column=7, value=fg.get("numero_comprobante", ""))
        ws.cell(row=row_idx, column=8, value=fg.get("medio_pago", ""))
        ws.cell(row=row_idx, column=9, value=cuit_val)
        ws.cell(row=row_idx, column=10, value=tipo_doc_val)
        ws.cell(row=row_idx, column=11, value=r.estrategia)
        ws.cell(row=row_idx, column=12, value=r.estado)

    wb.save(ruta_destino)
    wb.close()


def _comparar_con_gabarito(resultados, ruta_gabarito: Path):
    wb = openpyxl.load_workbook(ruta_gabarito, data_only=True)
    ws = wb.active

    headers = {str(ws.cell(row=1, column=c).value).strip().lower(): c for c in range(1, ws.max_column + 1)}
    col_arch = headers.get("archivo", 1)
    col_fecha = headers.get("fecha")
    col_monto = headers.get("monto")
    col_prov = headers.get("proveedor")
    col_cat = headers.get("categoría") or headers.get("categoria")

    verdad = {}
    for r in range(2, ws.max_row + 1):
        arch_nombre = ws.cell(row=r, column=col_arch).value
        if arch_nombre:
            verdad[str(arch_nombre).strip()] = {
                "fecha": str(ws.cell(row=r, column=col_fecha).value).strip() if col_fecha and ws.cell(row=r, column=col_fecha).value else None,
                "monto": float(ws.cell(row=r, column=col_monto).value) if col_monto and ws.cell(row=r, column=col_monto).value else None,
                "proveedor": str(ws.cell(row=r, column=col_prov).value).strip() if col_prov and ws.cell(row=r, column=col_prov).value else None,
                "categoria": str(ws.cell(row=r, column=col_cat).value).strip() if col_cat and ws.cell(row=r, column=col_cat).value else None,
            }
    wb.close()

    aciertos_fecha = 0
    aciertos_monto = 0
    aciertos_prov = 0
    total_comparables = 0

    for r in resultados:
        arch_name = r.archivo.name
        if arch_name not in verdad:
            continue
        v = verdad[arch_name]
        total_comparables += 1
        fg = r.fila_gastos or {}

        if v["fecha"] and formatear_fecha_fudo(fg.get("fecha")) == formatear_fecha_fudo(v["fecha"]):
            aciertos_fecha += 1
        if v["monto"] is not None and abs(float(fg.get("monto", 0.0)) - v["monto"]) <= 1.0:
            aciertos_monto += 1
        if v["proveedor"] and fg.get("proveedor", "").strip().lower() == v["proveedor"].lower():
            aciertos_prov += 1

    if total_comparables > 0:
        print(f"\n[*] Métricas de precisión ({total_comparables} comprobantes en gabarito):")
        print(f"    - Fecha exacta:      {aciertos_fecha}/{total_comparables} ({aciertos_fecha / total_comparables:.1%})")
        print(f"    - Monto (tol. $1):   {aciertos_monto}/{total_comparables} ({aciertos_monto / total_comparables:.1%})")
        print(f"    - Proveedor:         {aciertos_prov}/{total_comparables} ({aciertos_prov / total_comparables:.1%})")
    else:
        print("    [!] Los nombres de archivo del lote no coinciden con los del gabarito.")
